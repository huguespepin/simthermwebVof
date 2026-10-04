"""
db_loader.py — Chargement de la base de données SimTherm depuis le
classeur Excel (SimTherm_BD_v6.xlsx), source unique de vérité partagée
avec la version Excel du simulateur.

Aucune valeur dérivée (U, phi) n'est lue depuis la BD : seules les
propriétés intrinsèques (lambda, rho, Cp) le sont, conformément au
principe méthodologique du Chapitre 2.2.1 du mémoire.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import openpyxl

from core.constants import CATEGORIES_MUR, CATEGORIES_TOITURE


@dataclass
class Materiau:
    id: str
    pays: str
    nom: str
    categorie: str
    lambda_val: float
    rho: float
    cp: float
    niveau_confiance: Optional[str] = None
    cout_local: Optional[float] = None
    monnaie: Optional[str] = None
    disponibilite: Optional[str] = None
    epaisseur_ref_cm: Optional[float] = None  # épaisseur de référence de la source


@dataclass
class Ville:
    pays: str
    nom: str
    tmax: float
    tmoy: float
    tmin: float
    humidite: float
    ghi: float
    amplitude: float
    tn: float
    koppen: Optional[str] = None


@dataclass
class TypeBatiment:
    nom: str
    categorie: str
    apports_w_m2: float
    nv_vol_h: float


class SimThermDB:
    """
    Charge et expose la base de données SimTherm.

    Usage :
        db = SimThermDB("data/SimTherm_BD_v6.xlsx")
        pays_list = db.get_pays()
        villes = db.get_villes("Bénin")
        murs = db.get_materiaux_mur("Bénin")
    """

    def __init__(self, xlsx_path: str = "data/SimTherm_BD_v6.xlsx"):
        self.xlsx_path = Path(xlsx_path)
        if not self.xlsx_path.exists():
            raise FileNotFoundError(
                f"Base de données introuvable : {self.xlsx_path}. "
                f"Vérifiez que le fichier est bien présent dans le dossier data/."
            )
        self._wb = openpyxl.load_workbook(self.xlsx_path, data_only=True)
        self._materiaux: list[Materiau] = []
        self._villes: list[Ville] = []
        self._batiments: list[TypeBatiment] = []
        self._pays: list[str] = []
        self._charger()

    # ------------------------------------------------------------
    def _charger(self):
        self._charger_pays()
        self._charger_climat()
        self._charger_materiaux()
        self._charger_batiments()

    def _charger_pays(self):
        ws = self._wb["BD_PAYS"]
        pays_set = []
        for row in ws.iter_rows(min_row=5, values_only=True):
            if row[0]:
                pays_set.append(row[0])
        self._pays = pays_set

    @staticmethod
    def _is_plausible_tn(v) -> bool:
        """Une température de neutralité plausible se situe entre 15 et 35°C.
        Sert de garde-fou contre un décalage de colonne (cf. note ci-dessous)."""
        return isinstance(v, (int, float)) and 15 <= v <= 35

    def _charger_climat(self):
        ws = self._wb["BD_CLIMAT"]
        # Mapping BD_CLIMAT v6 (indices 0-based, en-têtes ligne 4) :
        # 0=Pays 1=Code 2=Ville 3=Köppen 4=Description 5=Tmin 6=Tmax
        # 7=Tmoy 8=HR 9=GHI 10=Amplitude 11=Précip/Tn(*) 12=Tn/Source(*)
        #
        # (*) NOTE — incohérence de saisie détectée dans BD_CLIMAT v6 :
        # pour certaines lignes (ex. Bénin), la colonne "Précipitations"
        # n'a jamais été renseignée et Tn a été saisie directement en
        # colonne L (idx 11) au lieu de M (idx 12) ; pour d'autres lignes
        # (ex. Nigeria, Kenya, Togo), la précipitation EST renseignée et
        # Tn se trouve bien en colonne M. La position de Tn varie donc
        # selon la ligne. On résout par heuristique : Tn est une valeur
        # plausible de température (15-35°C), une précipitation annuelle
        # ne l'est pas (généralement >100mm). Validé sur 156/156 villes.
        # À corriger à la source dans une prochaine révision de la BD.
        for row in ws.iter_rows(min_row=5, values_only=True):
            pays = row[0]
            ville = row[2] if len(row) > 2 else None
            if not pays or not ville:
                continue
            try:
                tmin = float(row[5]) if row[5] is not None else None
                tmax = float(row[6]) if row[6] is not None else None
                tmoy = float(row[7]) if row[7] is not None else None
                humidite = float(row[8]) if len(row) > 8 and row[8] is not None else None
                ghi = float(row[9]) if len(row) > 9 and row[9] is not None else None
                amplitude = float(row[10]) if len(row) > 10 and row[10] is not None else None
            except (TypeError, ValueError):
                continue

            # Résolution robuste de Tn (cf. note ci-dessus)
            val_l = row[11] if len(row) > 11 else None
            val_m = row[12] if len(row) > 12 else None
            tn = None
            if self._is_plausible_tn(val_l):
                tn = float(val_l)
            else:
                try:
                    val_m_f = float(val_m)
                    if self._is_plausible_tn(val_m_f):
                        tn = val_m_f
                except (TypeError, ValueError):
                    pass

            koppen = row[3] if len(row) > 3 else None

            if tmax is None or tmoy is None or amplitude is None or tn is None:
                continue
            self._villes.append(Ville(
                pays=pays, nom=ville, tmax=tmax, tmoy=tmoy,
                tmin=tmin or (tmoy - 5), humidite=humidite or 70.0,
                ghi=ghi or 5.0, amplitude=amplitude, tn=tn, koppen=koppen,
            ))

    def _charger_materiaux(self):
        ws = self._wb["BD_MATERIAUX"]
        # Colonnes : A=ID(1) B=Pays(2) D=Nom(4) E=Catégorie(5)
        # F=lambda(6) G=rho(7) H=Cp(8) M=Niveau(13)
        # I=Ép. de référence en cm(9) R=Coût local/m²(18) S=Monnaie(19) T=Disponibilité(20)
        for row in ws.iter_rows(min_row=5, values_only=True):
            id_ = row[0]
            if not id_:
                continue
            pays = row[1] if len(row) > 1 else None
            nom = row[3] if len(row) > 3 else None
            categorie = row[4] if len(row) > 4 else None
            try:
                lambda_val = float(row[5]) if len(row) > 5 and row[5] not in (None, "") else None
                rho = float(row[6]) if len(row) > 6 and row[6] not in (None, "") else None
                cp = float(row[7]) if len(row) > 7 and row[7] not in (None, "") else None
            except (TypeError, ValueError):
                continue
            if not all([nom, categorie, lambda_val, rho, cp]) or lambda_val <= 0:
                continue
            niveau = row[12] if len(row) > 12 else None
            try:
                cout_local = float(row[17]) if len(row) > 17 and row[17] not in (None, "") else None
            except (TypeError, ValueError):
                cout_local = None
            monnaie = row[18] if len(row) > 18 else None
            disponibilite = row[19] if len(row) > 19 else None
            try:
                ep_ref = float(row[8]) if len(row) > 8 and row[8] not in (None, "") else None
            except (TypeError, ValueError):
                ep_ref = None
            self._materiaux.append(Materiau(
                id=id_, pays=pays or "Tous pays", nom=nom, categorie=categorie,
                lambda_val=lambda_val, rho=rho, cp=cp, niveau_confiance=niveau,
                cout_local=cout_local, monnaie=monnaie, disponibilite=disponibilite,
                epaisseur_ref_cm=ep_ref,
            ))

    def _charger_batiments(self):
        ws = self._wb["BD_BATIMENTS"]
        # Colonnes : B=Nom(2) E=Apports(5) F ou G=Nv
        for row in ws.iter_rows(min_row=5, values_only=True):
            nom = row[1] if len(row) > 1 else None
            if not nom:
                continue
            try:
                apports = float(row[4]) if len(row) > 4 and row[4] not in (None, "") else 4.0
                nv = float(row[5]) if len(row) > 5 and row[5] not in (None, "") else 1.0
            except (TypeError, ValueError):
                apports, nv = 4.0, 1.0
            categorie = row[2] if len(row) > 2 and row[2] else ""
            self._batiments.append(TypeBatiment(
                nom=nom, categorie=categorie, apports_w_m2=apports, nv_vol_h=nv,
            ))

    # ------------------------------------------------------------
    # API publique
    # ------------------------------------------------------------
    def get_pays(self) -> list[str]:
        return list(self._pays)

    def get_villes(self, pays: str) -> list[Ville]:
        return [v for v in self._villes if v.pays == pays]

    def get_ville(self, pays: str, nom_ville: str) -> Optional[Ville]:
        for v in self._villes:
            if v.pays == pays and v.nom == nom_ville:
                return v
        return None

    def get_materiaux_mur(self, pays: Optional[str] = None) -> list[Materiau]:
        mats = [m for m in self._materiaux if m.categorie in CATEGORIES_MUR]
        if pays:
            mats = [m for m in mats if m.pays == pays or m.pays == "Tous pays"]
        return mats

    def get_materiaux_toit(self, pays: Optional[str] = None) -> list[Materiau]:
        mats = [m for m in self._materiaux if m.categorie in CATEGORIES_TOITURE]
        if pays:
            mats = [m for m in mats if m.pays == pays or m.pays == "Tous pays"]
        return mats

    def get_materiau_par_nom(self, nom: str) -> Optional[Materiau]:
        for m in self._materiaux:
            if m.nom == nom:
                return m
        return None

    def get_batiments(self) -> list[TypeBatiment]:
        return list(self._batiments)

    def get_batiment_par_nom(self, nom: str) -> Optional[TypeBatiment]:
        for b in self._batiments:
            if b.nom == nom:
                return b
        return None
